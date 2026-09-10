"""Pobiera adresy plakatów z TMDB i zapisuje je do arkusza.

Wejście:  filmy.xlsx, arkusz "Filmy", kolumny "Tytuł" i "Rok".
Wyjście:  plakaty.xlsx, kolumny "Tytuł", "Rok", "FilmId", "Plakat".

To jedyny moment, w którym cokolwiek rozmawia z TMDB. Aplikacja Power Apps
nie odpytuje API — dostaje gotową listę adresów wklejoną do App.OnStart.
Skrypt uruchamiasz ponownie tylko wtedy, gdy dopiszesz filmy do bazy.

Klucz API czytany jest ze zmiennej środowiskowej, nigdy z kodu:

    export TMDB_API_KEY="twoj_klucz"        # klucz v3
    python pobierz_plakaty.py

Klucz zakładasz na https://www.themoviedb.org/settings/api (darmowy).

Użycie:
    python pobierz_plakaty.py [--zrodlo filmy.xlsx] [--wynik plakaty.xlsx]
                              [--szerokosc w300]
"""

from __future__ import annotations

import argparse
import os
import sys
import time
import urllib.parse
import urllib.request
import json

import openpyxl

API = "https://api.themoviedb.org/3/search/movie"
OBRAZKI = "https://image.tmdb.org/t/p/"

# TMDB potrafi zwrócić kilka trafień na ten sam tytuł. Kolejność prób:
# najpierw z rokiem produkcji, potem bez — polskie filmy przedwojenne
# bywają w bazie z rozjechaną datą.
PROBY = [True, False]


def klucz() -> str:
    k = os.environ.get("TMDB_API_KEY", "").strip()
    if not k:
        sys.exit(
            "Brak klucza. Ustaw zmienną środowiskową TMDB_API_KEY:\n"
            '    export TMDB_API_KEY="twoj_klucz"'
        )
    return k


def zapytaj(tytul: str, rok: int, k: str, z_rokiem: bool) -> dict | None:
    """Zwraca pierwszy wynik wyszukiwania albo None."""
    params = {
        "api_key": k,
        "query": tytul,
        "language": "pl-PL",
        "include_adult": "false",
    }
    if z_rokiem:
        params["year"] = str(rok)

    url = f"{API}?{urllib.parse.urlencode(params)}"
    with urllib.request.urlopen(url, timeout=20) as r:
        dane = json.load(r)

    wyniki = dane.get("results") or []
    return wyniki[0] if wyniki else None


def plakat(tytul: str, rok: int, k: str, szerokosc: str) -> str:
    """Pełny adres plakatu albo pusty string, gdy TMDB nic nie ma."""
    for z_rokiem in PROBY:
        try:
            trafienie = zapytaj(tytul, rok, k, z_rokiem)
        except Exception as e:                      # sieć, limit, cokolwiek
            print(f"  ! błąd zapytania ({e}) — próbuję dalej", file=sys.stderr)
            time.sleep(2)
            continue
        if trafienie and trafienie.get("poster_path"):
            return OBRAZKI + szerokosc + trafienie["poster_path"]
    return ""


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--zrodlo", default="filmy.xlsx")
    ap.add_argument("--wynik", default="plakaty.xlsx")
    ap.add_argument("--szerokosc", default="w300",
                    help="rozmiar plakatu wg TMDB: w185, w300, w500, original")
    a = ap.parse_args()

    k = klucz()

    ws = openpyxl.load_workbook(a.zrodlo)["Filmy"]
    naglowki = [str(c.value).strip() if c.value else "" for c in ws[1]]
    i_tytul, i_rok = naglowki.index("Tytuł"), naglowki.index("Rok")

    out = openpyxl.Workbook()
    wo = out.active
    wo.title = "Plakaty"
    wo.append(["Tytuł", "Rok", "FilmId", "Plakat"])

    brakujace = []
    for wiersz in ws.iter_rows(min_row=2, values_only=True):
        tytul = wiersz[i_tytul]
        if not tytul:
            continue                                # pusty wiersz na dole arkusza
        tytul, rok = str(tytul).strip(), int(wiersz[i_rok])
        film_id = f"{tytul}|{rok}"

        url = plakat(tytul, rok, k, a.szerokosc)
        if not url:
            brakujace.append(film_id)
            print(f"  – bez plakatu: {film_id}")
        else:
            print(f"  ✓ {film_id}")

        wo.append([tytul, rok, film_id, url])
        time.sleep(0.25)                            # nie zajeżdżamy API

    for kol, szer in zip("ABCD", (34, 8, 38, 74)):
        wo.column_dimensions[kol].width = szer
    wo.freeze_panes = "A2"
    out.save(a.wynik)

    print(f"\nZapisano {a.wynik}: {wo.max_row - 1} wierszy.")
    if brakujace:
        print(f"Bez plakatu ({len(brakujace)}): " + ", ".join(brakujace))
        print("Uzupełnij te adresy ręcznie w arkuszu albo zostaw puste —")
        print("karta filmu poradzi sobie bez plakatu.")


if __name__ == "__main__":
    main()
